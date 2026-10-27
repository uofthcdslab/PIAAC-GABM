#!/usr/bin/env python3
"""Create compact persona descriptions from the weighted PIAAC sample."""

from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = PROJECT_ROOT / "data" / "piaac_cycle2_weighted_sample.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "piaac_personas.json"

CATEGORICAL_FIELDS = {
    "AGEG10LFS": "age",
    "GENDER_R": "gender",
    "IMGENC2": "immigration background",
    "LNG_BQ": "questionnaire language",
    "NATIVELANG": "language alignment",
    "EDA_ISCED1_C": "educational attainment",
    "C2_D05": "employment status",
    "ISCO1C": "occupation",
    "D2_Q04": "employment type",
    "D2_Q13": "job satisfaction",
    "MONTHLYINCPR": "monthly income percentile",
    "PARTNER": "household",
    "I2_Q01b": "social trust (0-10)",
    "I2_Q03": "health",
}

SKILL_SOURCE_FIELDS = {
    "G2_Q04": "ever used a digital device",
    "G2_Q05c": "tablet use frequency",
    "G2_Q06a": "email use in everyday life",
    "G2_Q06b": "internet information use in everyday life",
    "G2_Q06c": "online transactions in everyday life",
    "G2_Q06d": "digital content creation in everyday life",
    "G2_Q06e": "digital personal-life management",
    "F2_Q04": "computer experience at work",
    "F2_Q05a": "online communication at work",
    "F2_Q05c": "online information use at work",
    "F2_Q05d": "digital document creation at work",
    "F2_Q05e": "specialized software use at work",
    "F2_Q05f": "programming-language use at work",
    "H2_Q09a": "learning from coworkers",
    "H2_Q09b": "learning by doing",
    "H2_Q09c": "keeping up with new products or services",
    "FE12C2": "formal education in the last 12 months",
    "NFE12C2": "non-formal education in the last 12 months",
    "NFE12JRC2": "job-related non-formal education",
    "NFE12NJRC2": "non-job-related non-formal education",
}

# These values are retained numerically because uncertainty/calibration matters.
EXACT_NUMERIC_FIELDS = {
    "YRSQUALC2": "years of formal education",
    "ICTHOMEC2_SE": "digital technology use at home SE",
    "ICTWORKC2_SE": "digital technology use at work SE",
    "LEARNATWORKC2_SE": "learning at work SE",
    "EffortP001S": "general effort calibration score",
    "EffortP002S": "effort control score",
}

# Numeric fields are converted to relative low/typical/high descriptions using
# the 25th and 75th percentiles of this sample.
NUMERIC_FIELDS = {
    "ICTHOMEC2": "digital technology use at home",
    "READHOMEC2_T1": "reading use at home",
    "NUMHOMEC2": "numeracy use at home",
    "AGRE": "agreeableness",
    "CONS": "conscientiousness",
    "EMOS": "emotional stability",
    "EXTR": "extraversion",
    "OPEM": "open-mindedness",
    "OPEM_CURI": "intellectual curiosity",
    "EMOS_ANXI": "anxiety (weak general-affect context)",
}

WORK_NUMERIC_FIELDS = {
    "ICTWORKC2": "digital technology use at work",
    "READWORKC2_T1": "reading use at work",
    "NUMWORKC2": "numeracy use at work",
    "LEARNATWORKC2": "learning at work",
    "TASKDISCC2_T1": "task discretion at work",
}

MISSING_PREFIXES = (".",)
MISSING_TEXT = {"", "none", "nan", "not applicable", "valid skip"}


def variable_name(header: Any) -> str:
    """Extract the PIAAC variable name from a readable workbook header."""
    return str(header).split(" — ", 1)[0].strip()


def clean_label(value: Any) -> str | None:
    """Remove numeric codes and normalize PIAAC missing-value labels."""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None

    text = str(value).strip()
    if not text or text.lower() in MISSING_TEXT or text.startswith(MISSING_PREFIXES):
        return None

    if " — " in text:
        code, label = text.split(" — ", 1)
        if code.startswith(MISSING_PREFIXES) or label.strip().lower() in MISSING_TEXT:
            return None
        return label.strip()
    return text


def numeric_value(value: Any) -> float | None:
    cleaned = clean_label(value)
    if cleaned is None:
        return None
    try:
        return float(cleaned)
    except (TypeError, ValueError):
        return None


def percentile(values: list[float], probability: float) -> float:
    """Return a linearly interpolated percentile without extra dependencies."""
    ordered = sorted(values)
    if not ordered:
        raise ValueError("Cannot calculate a percentile from an empty list.")
    position = (len(ordered) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    fraction = position - lower
    return ordered[lower] + (ordered[upper] - ordered[lower]) * fraction


def relative_level(value: float, lower: float, upper: float) -> str:
    if value <= lower:
        return "lower"
    if value >= upper:
        return "higher"
    return "typical"


def format_number(value: float) -> str:
    """Keep scores readable while retaining useful precision."""
    return f"{value:.3f}".rstrip("0").rstrip(".")


def make_persona_id(sequence_id: Any, occurrences: Counter[str], row_number: int) -> str:
    base = clean_label(sequence_id) or f"row-{row_number}"
    base = f"PIAAC-{base}"
    occurrences[base] += 1
    return base if occurrences[base] == 1 else f"{base}-{occurrences[base]}"


THEME_ORDER = (
    "demography",
    "education",
    "employment",
    "household_and_economic_context",
    "core_skills",
    "digital_skills",
    "learning_and_training",
    "personality",
    "social_and_health",
    "calibration",
)


def build_themes(
    row: tuple[Any, ...],
    indexes: dict[str, int],
    thresholds: dict[str, tuple[float, float]],
) -> dict[str, str]:
    themes: dict[str, list[str]] = {theme: [] for theme in THEME_ORDER}

    categorical_themes = {
        "AGEG10LFS": "demography",
        "GENDER_R": "demography",
        "IMGENC2": "demography",
        "LNG_BQ": "demography",
        "NATIVELANG": "demography",
        "EDA_ISCED1_C": "education",
        "C2_D05": "employment",
        "ISCO1C": "employment",
        "D2_Q04": "employment",
        "D2_Q13": "employment",
        "MONTHLYINCPR": "household_and_economic_context",
        "PARTNER": "household_and_economic_context",
        "I2_Q01b": "social_and_health",
        "I2_Q03": "social_and_health",
    }

    for variable, label in CATEGORICAL_FIELDS.items():
        value = clean_label(row[indexes[variable]])
        if value is not None:
            themes[categorical_themes[variable]].append(f"{label}: {value}")

    source_themes = {
        variable: "digital_skills"
        for variable in SKILL_SOURCE_FIELDS
        if variable.startswith(("G2_", "F2_"))
    }
    source_themes.update(
        {
            variable: "learning_and_training"
            for variable in SKILL_SOURCE_FIELDS
            if variable.startswith(("H2_", "FE", "NFE"))
        }
    )

    for variable, label in SKILL_SOURCE_FIELDS.items():
        value = clean_label(row[indexes[variable]])
        if value is not None:
            themes[source_themes[variable]].append(f"{label}: {value}")

    exact_themes = {
        "YRSQUALC2": "education",
        "ICTHOMEC2_SE": "digital_skills",
        "ICTWORKC2_SE": "digital_skills",
        "LEARNATWORKC2_SE": "learning_and_training",
        "EffortP001S": "calibration",
        "EffortP002S": "calibration",
    }

    for variable, label in EXACT_NUMERIC_FIELDS.items():
        value = numeric_value(row[indexes[variable]])
        if value is not None:
            themes[exact_themes[variable]].append(f"{label}: {format_number(value)}")

    relative_themes = {
        "ICTHOMEC2": "digital_skills",
        "ICTWORKC2": "digital_skills",
        "READHOMEC2_T1": "core_skills",
        "READWORKC2_T1": "core_skills",
        "NUMHOMEC2": "core_skills",
        "NUMWORKC2": "core_skills",
        "LEARNATWORKC2": "learning_and_training",
        "TASKDISCC2_T1": "employment",
        "AGRE": "personality",
        "CONS": "personality",
        "EMOS": "personality",
        "EXTR": "personality",
        "OPEM": "personality",
        "OPEM_CURI": "personality",
        "EMOS_ANXI": "personality",
    }

    for variable, label in {**NUMERIC_FIELDS, **WORK_NUMERIC_FIELDS}.items():
        value = numeric_value(row[indexes[variable]])
        if value is not None:
            lower, upper = thresholds[variable]
            themes[relative_themes[variable]].append(
                f"{label}: {relative_level(value, lower, upper)} relative to the sample"
            )

    return {
        theme: "; ".join(facts) + "."
        for theme, facts in themes.items()
        if facts
    }


def create_personas(input_path: Path, output_path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    indexes = {variable_name(header): index for index, header in enumerate(headers)}

    required = {
        "SEQID",
        *CATEGORICAL_FIELDS,
        *SKILL_SOURCE_FIELDS,
        *EXACT_NUMERIC_FIELDS,
        *NUMERIC_FIELDS,
        *WORK_NUMERIC_FIELDS,
    }
    missing = sorted(required - indexes.keys())
    if missing:
        workbook.close()
        raise KeyError(f"Required columns are missing: {', '.join(missing)}")

    data_rows = list(rows)
    workbook.close()

    thresholds: dict[str, tuple[float, float]] = {}
    for variable in (*NUMERIC_FIELDS, *WORK_NUMERIC_FIELDS):
        values = [
            number
            for row in data_rows
            if (number := numeric_value(row[indexes[variable]])) is not None
        ]
        thresholds[variable] = (percentile(values, 0.25), percentile(values, 0.75))

    occurrences: Counter[str] = Counter()
    personas: dict[str, dict[str, str]] = {}
    for excel_row, row in enumerate(data_rows, start=2):
        persona_id = make_persona_id(row[indexes["SEQID"]], occurrences, excel_row)
        personas[persona_id] = build_themes(row, indexes, thresholds)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(personas, ensure_ascii=False, indent=4),
        encoding="utf-8",
    )
    return personas


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    personas = create_personas(args.input, args.output)
    print(f"Created {len(personas):,} personas at {args.output}")


if __name__ == "__main__":
    main()
