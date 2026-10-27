import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = PROJECT_ROOT / "outputs" / "parsed_responses.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "outputs" / "parsed_responses_without_3.xlsx"
JSON_OUTPUT = PROJECT_ROOT / "outputs" / "policy_positions.json"
EXCLUDED_RATINGS = {3}


source = load_workbook(DEFAULT_INPUT, read_only=True, data_only=True)
output = Workbook(write_only=True)
policy_positions = {}

for source_sheet in source.worksheets:
    policy_positions[source_sheet.title] = {"against": [], "for": []}
    output_sheet = output.create_sheet(source_sheet.title)
    rows = source_sheet.iter_rows(values_only=True)
    header = next(rows)
    output_sheet.append(header)
    agent_id_index = header.index("Agent ID")
    rating_index = header.index("Rating")

    for row in rows:
        rating = row[rating_index]
        if rating not in EXCLUDED_RATINGS:
            output_sheet.append(row)
        if rating in {1, 2}:
            policy_positions[source_sheet.title]["against"].append(row[agent_id_index])
        elif rating in {4, 5}:
            policy_positions[source_sheet.title]["for"].append(row[agent_id_index])

output.save(DEFAULT_OUTPUT)
source.close()
with open(JSON_OUTPUT, "w", encoding="utf-8") as file:
    json.dump(policy_positions, file, ensure_ascii=False, indent=2)

print(f"Saved filtered workbook to {DEFAULT_OUTPUT}")
print(f"Saved policy positions to {JSON_OUTPUT}")
for policy, positions in policy_positions.items():
    print(
        f"{policy}: "
        f"against={len(positions['against'])}, "
        f"for={len(positions['for'])}"
    )
