#!/usr/bin/env python3
"""Split the weighted PIAAC sample into 10 balanced subsamples of 320 rows."""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import xlsxwriter


INPUT_FILE = Path("data/piaac_cycle2_weighted_sample.xlsx")
OUTPUT_FILE = Path(
    "outputs/01a0b9c2-864b-7540-8d96-b87786b0c594/"
    "piaac_cycle2_balanced_subsamples.xlsx"
)

BALANCE_VARIABLES = [
    "GENDER_R",
    "AGEG10LFS",
    "LNG_BQ",
    "EDU_BQDI_C",
]

NUMBER_OF_SUBSAMPLES = 10
ROWS_PER_SUBSAMPLE = 320
RANDOM_SEED = 20260921
OPTIMIZATION_STEPS = 500_000


def variable_name(header: object) -> str:
    return str(header).split(" — ", 1)[0]


def load_balancing_data() -> tuple[list[object], list[list[object]]]:
    workbook = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    names = [variable_name(header) for header in headers]
    indexes = [names.index(variable) for variable in BALANCE_VARIABLES]

    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append([row[index] if row[index] is not None else "Missing" for index in indexes])
    workbook.close()
    return headers, rows


def one_hot_matrix(rows: list[list[object]]) -> tuple[np.ndarray, list[tuple[str, object]]]:
    columns = []
    labels = []
    for variable_index, variable in enumerate(BALANCE_VARIABLES):
        categories = sorted({row[variable_index] for row in rows}, key=str)
        for category in categories:
            columns.append([1 if row[variable_index] == category else 0 for row in rows])
            labels.append((variable, category))
    return np.asarray(columns, dtype=np.int8).T, labels


def objective(counts: np.ndarray, target: np.ndarray) -> float:
    scale = np.maximum(target, 1.0)
    return float(np.sum((counts - target) ** 2 / scale))


def assign_subsamples(matrix: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    row_count = matrix.shape[0]
    expected = NUMBER_OF_SUBSAMPLES * ROWS_PER_SUBSAMPLE
    if row_count != expected:
        raise ValueError(f"Expected {expected} rows, found {row_count}.")

    rng = np.random.default_rng(RANDOM_SEED)
    assignments = np.repeat(np.arange(NUMBER_OF_SUBSAMPLES), ROWS_PER_SUBSAMPLE)
    rng.shuffle(assignments)

    counts = np.zeros((NUMBER_OF_SUBSAMPLES, matrix.shape[1]), dtype=np.int32)
    for group in range(NUMBER_OF_SUBSAMPLES):
        counts[group] = matrix[assignments == group].sum(axis=0)

    target = matrix.sum(axis=0) / NUMBER_OF_SUBSAMPLES
    scale = np.maximum(target, 1.0)

    for step in range(OPTIMIZATION_STEPS):
        first = rng.integers(row_count)
        second = rng.integers(row_count)
        group_a = assignments[first]
        group_b = assignments[second]
        if group_a == group_b:
            continue

        old_a = counts[group_a]
        old_b = counts[group_b]
        new_a = old_a - matrix[first] + matrix[second]
        new_b = old_b - matrix[second] + matrix[first]

        old_score = np.sum((old_a - target) ** 2 / scale) + np.sum(
            (old_b - target) ** 2 / scale
        )
        new_score = np.sum((new_a - target) ** 2 / scale) + np.sum(
            (new_b - target) ** 2 / scale
        )

        if new_score < old_score:
            assignments[first], assignments[second] = group_b, group_a
            counts[group_a] = new_a
            counts[group_b] = new_b

    return assignments, counts, target


def write_output(headers: list[object], assignments: np.ndarray) -> None:
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    source_workbook = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    source_sheet = source_workbook.active

    output_workbook = xlsxwriter.Workbook(
        OUTPUT_FILE,
        {
            "constant_memory": True,
            "strings_to_formulas": False,
            "strings_to_urls": False,
        },
    )
    output_sheet = output_workbook.add_worksheet("Balanced subsamples")
    output_sheet.write_row(0, 0, ["SUBSAMPLE_ID", *headers])

    output_row = 1
    for assignment, row in zip(
        assignments, source_sheet.iter_rows(min_row=2, values_only=True)
    ):
        output_sheet.write_row(output_row, 0, [int(assignment) + 1, *row])
        output_row += 1

    output_sheet.autofilter(0, 0, output_row - 1, len(headers))
    output_sheet.freeze_panes(1, 1)
    output_workbook.close()
    source_workbook.close()


def main() -> None:
    headers, balancing_rows = load_balancing_data()
    matrix, labels = one_hot_matrix(balancing_rows)
    assignments, counts, target = assign_subsamples(matrix)
    write_output(headers, assignments)

    print(f"Created {OUTPUT_FILE}")
    print(f"Rows per subsample: {np.bincount(assignments).tolist()}")
    print(f"Balance objective: {objective(counts, target):.6f}")
    print(f"Largest category-count deviation: {np.max(np.abs(counts - target)):.3f}")
    for column, (variable, category) in enumerate(labels):
        values = counts[:, column].tolist()
        print(f"{variable} | {category}: target={target[column]:.1f}, samples={values}")


if __name__ == "__main__":
    main()
